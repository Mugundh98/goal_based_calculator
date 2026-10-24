"""Approximate preview of the presentation, for checking layout without PowerPoint.

    report/.venv/Scripts/python.exe report/preview_slides.py build/GoalWise_Presentation.pptx

Draws every slide's shapes, pictures and text (wrapped with the real Calibri / Cambria metrics)
to build/slides/slide-NN.png and reports text that would overflow its box. Charts are drawn as a
labelled placeholder. It is a check for overflow and overlap, not a faithful renderer.
"""

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

SCALE = 120  # pixels per inch
FONTS = Path("C:/Windows/Fonts")
THEME = {"dk1": "1F2933", "lt1": "FFFFFF", "dk2": "0B3B32", "lt2": "E7F5EF", "accent1": "059669",
         "accent2": "F59E0B", "accent3": "0E7490", "accent4": "64748B", "accent5": "10B981",
         "accent6": "B45309", "tx1": "1F2933", "tx2": "0B3B32", "bg1": "FFFFFF", "bg2": "E7F5EF"}


def font(face: str, size_pt: float, bold: bool) -> ImageFont.FreeTypeFont:
    px = max(int(size_pt / 72 * SCALE), 6)
    if face.startswith("Cambria"):
        return ImageFont.truetype(str(FONTS / ("cambriab.ttf" if bold else "cambria.ttc")), px)
    return ImageFont.truetype(str(FONTS / ("calibrib.ttf" if bold else "calibri.ttf")), px)


def color_of(fill_or_font, default="1F2933") -> str:
    try:
        c = fill_or_font.color
        if c.type is None:
            return default
        if c.type == 1:  # RGB
            return str(c.rgb)
        name = str(c.theme_color).split(".")[-1].split(" ")[0].lower()
        mapping = {"text_1": "tx1", "text_2": "tx2", "background_1": "bg1", "background_2": "bg2",
                   "dark_1": "dk1", "light_1": "lt1", "dark_2": "dk2", "light_2": "lt2"}
        return THEME.get(mapping.get(name, name), default)
    except Exception:
        return default


def px(emu: int) -> int:
    return int(Emu(emu).inches * SCALE)


def wrap(text: str, f: ImageFont.FreeTypeFont, width: int) -> list[str]:
    lines = []
    for para in text.split("\n"):
        words, line = para.split(" "), ""
        for w in words:
            trial = (line + " " + w).strip()
            if f.getlength(trial) <= width or not line:
                line = trial
            else:
                lines.append(line)
                line = w
        lines.append(line)
    return lines


def main() -> None:
    deck = Path(sys.argv[1])
    out = deck.parent / "slides"
    out.mkdir(exist_ok=True)
    prs = Presentation(str(deck))
    W, H = px(prs.slide_width), px(prs.slide_height)
    problems = []
    for n, slide in enumerate(prs.slides, start=1):
        layout_name = slide.slide_layout.name
        dark = layout_name in ("TITLE", "DARK")
        img = Image.new("RGB", (W, H), "#" + (THEME["dk2"] if dark else "FFFFFF"))
        draw = ImageDraw.Draw(img)
        for shape in slide.shapes:
            x, y, w, h = px(shape.left or 0), px(shape.top or 0), px(shape.width or 0), px(shape.height or 0)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    pic = Image.open(__import__("io").BytesIO(shape.image.blob)).convert("RGB")
                    pic.thumbnail((w, h))
                    img.paste(pic, (x + (w - pic.width) // 2, y + (h - pic.height) // 2))
                except Exception:
                    draw.rectangle([x, y, x + w, y + h], outline="red")
                continue
            if shape.has_chart if hasattr(shape, "has_chart") else False:
                draw.rectangle([x, y, x + w, y + h], outline="#94A3B8", fill="#F8FAFC")
                draw.text((x + 10, y + 10), f"[chart: {shape.chart.chart_type}]", fill="#475569")
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not shape.has_text_frame or (
                    shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not shape.text_frame.text.strip()):
                fill = "#" + color_of(shape.fill.fore_color and shape.fill, "E7F5EF") if shape.fill.type == 1 else None
                kind = str(shape.auto_shape_type)
                if "OVAL" in kind:
                    draw.ellipse([x, y, x + w, y + h], fill=fill)
                else:
                    draw.rounded_rectangle([x, y, x + w, y + h], radius=10, fill=fill)
                continue
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.fill.type == 1:
                draw.rounded_rectangle([x, y, x + w, y + h], radius=10,
                                       fill="#" + color_of(shape.fill, "059669"))
            # Text: size, face and colour from the first run, else layout defaults.
            run = next((r for p in shape.text_frame.paragraphs for r in p.runs), None)
            is_title = shape.is_placeholder and shape.placeholder_format.type is not None and "TITLE" in str(shape.placeholder_format.type)
            size = run.font.size.pt if run is not None and run.font.size else (
                (44 if layout_name == "TITLE" else 32 if layout_name == "DARK" else 30) if is_title
                else (18 if layout_name == "TITLE" else 15))
            bold = bool(run.font.bold) if run is not None and run.font.bold is not None else is_title
            face = (run.font.name if run is not None and run.font.name else
                    ("Cambria" if is_title else "Calibri"))
            colour = color_of(run.font, "FFFFFF" if dark else "1F2933") if run is not None else (
                "FFFFFF" if dark else "0B3B32")
            f = font(face, size, bold)
            bullet_indent = 18 if any(p._p.pPr is not None and p._p.pPr.find(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None
                for p in shape.text_frame.paragraphs) else 0
            text = "\n".join(p.text for p in shape.text_frame.paragraphs)
            lines = wrap(text, f, max(w - bullet_indent, 10))
            line_h = int(size / 72 * SCALE * 1.2)
            gap = int(6 / 72 * SCALE) if bullet_indent else 0
            needed = len(lines) * line_h + gap * (len(shape.text_frame.paragraphs) - 1)
            ty = y
            anchor = str(shape.text_frame.vertical_anchor)
            if "MIDDLE" in anchor:
                ty = y + (h - needed) // 2
            for line in lines:
                draw.text((x + bullet_indent, ty), line, font=f, fill="#" + colour)
                ty += line_h
            if needed > h + 4:
                problems.append(f"slide {n}: text overflows by {(needed - h) / SCALE:.2f} in: {text[:60]!r}")
                draw.rectangle([x, y, x + w, y + h], outline="red", width=2)
            if x + w > W or y + h > H:
                problems.append(f"slide {n}: box runs off the slide: {text[:40]!r}")
        img.save(out / f"slide-{n:02d}.png")
    print(f"{len(prs.slides)} slides drawn to {out}")
    print("\n".join(problems) if problems else "no overflow found")


if __name__ == "__main__":
    main()
